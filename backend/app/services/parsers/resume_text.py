from pathlib import Path

from docx import Document
from pypdf import PdfReader


class UnsupportedResumeFileType(ValueError):
    pass


class ResumeTextExtractor:
    supported_extensions = {".pdf", ".docx", ".txt"}

    def extract_text(self, file_path: Path) -> str:
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(file_path)
        if suffix == ".docx":
            return self._extract_docx(file_path)
        if suffix == ".txt":
            return file_path.read_text(encoding="utf-8", errors="ignore")
        raise UnsupportedResumeFileType(f"暂不支持 {suffix or '未知'} 格式")

    def _extract_pdf(self, file_path: Path) -> str:
        reader = PdfReader(str(file_path))
        texts = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(text for text in texts if text.strip())

    def _extract_docx(self, file_path: Path) -> str:
        document = Document(str(file_path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        table_texts: list[str] = []
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    table_texts.append(" ".join(values))
        return "\n".join([*paragraphs, *table_texts])
